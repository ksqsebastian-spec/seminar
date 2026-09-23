from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
out=Path('public/material')
texts={
'03-schrankbrief.txt':'''EIN SCHRANK FÜR HADI – UNSER SEMINARPROJEKT

Hadi Teherani ist eine reale Person. Unser Betrieb Werkraum, die Schrankidee und alle unten genannten Maße und Kostenannahmen sind erfunden. Es gibt keine Bestellung, Anfrage oder Zusage von Hadi Teherani oder seinem Büro.

INTERNER ÜBUNGSAUFTRAG
Von: Werkraum-Team (fiktiv)
Wir möchten einen maßgefertigten Büroschrank als unverbindliche Idee für das Büro von Hadi Teherani vorbereiten. Recherchiere öffentlich belegbare Fakten und den offiziellen Bürokontakt. Erstelle daraus eine eigene Idee, eine Kalkulation und einen Angebotsentwurf.

UNSER ENTWURF
Vorgang: SCH-026
Anbieter: Werkraum – fiktiver Handwerksbetrieb
Bezeichnung: Maßgefertigter Büroschrank
Breite 120 cm, Höhe 180 cm, Tiefe 45 cm – ausschließlich unsere Übungsannahme
Zwei Türen, vier Einlegeböden, Oberfläche in Eiche-Optik als Vorschlag
Kein bestätigter Kundenwunsch. Kein geprüfter Konstruktions- oder Zuschnittplan.

KALKULATIONSANNAHMEN
8 Materialeinheiten/Platten zu je 45,00 EUR netto = 360,00 EUR
1 Beschlagsatz zu 120,00 EUR netto
8 Arbeitsstunden zu je 65,00 EUR netto
1 Transportpauschale zu 80,00 EUR netto
20 % Übungsaufschlag auf die Summe
19 % Umsatzsteuer als Rechenannahme für die Übung
Alle Startwerte sind fiktiv. Materialeinheiten sind eine vereinfachte Kalkulationsannahme, keine Stückliste für eine Konstruktion. Recherchierte Preise nur bei passender Einheit und eindeutiger Preisbasis übernehmen.

NOCH OFFEN
Tatsächlicher Bedarf, Maße vor Ort, gewünschte Innenaufteilung, Oberfläche, Lieferkosten und Liefertermin. Büroanschrift und öffentliche Kontaktadresse an einer offiziellen Quelle prüfen. Keine persönliche Adresse erraten.

DEINE DATEIEN
Projekt: Schrank für Hadi
OneDrive-Ordner: Seminar – Schrank für Hadi
ChatGPT soll den Ordner anlegen, die hochgeladene Quelle ablegen, eine bearbeitbare Word-Fassung erstellen und zusätzlich eine PDF erzeugen.
Das Angebot bleibt als unverbindlicher Seminarentwurf gekennzeichnet. Übungsmails an Kerims zuvor geprüfte Adresse speichern, nicht versenden.
''',
'09-recherche-und-browser.txt':'''ZWEI FÄHIGKEITEN – ZWEI AUFGABEN

1 INTERNETRECHERCHE
Wer ist Hadi Teherani? Finde drei kurze Fakten auf offiziellen Quellen. Suche die öffentlich genannte E-Mail-Adresse seines Designbüros. Prüfe, zu welcher Gesellschaft/Abteilung sie gehört. Eine Büro-Adresse ist keine persönliche E-Mail.
Start: https://www.haditeherani.com/
Offizielle Vita als mögliche weitere Quelle:
https://www.haditeherani.com/sites/default/files/2025-02/211208_HT%20Vita%20dt.pdf
Im Suchindex der offiziellen Vita ist design@haditeherani.com der Hadi Teherani Design GmbH zugeordnet (Recherche 23.09.2026). Vor Nutzung auf der zugänglichen Originalquelle erneut prüfen. Direkter Abruf kann durch die Website eingeschränkt sein. Keine Anfrage versenden.
Sichere die Fakten, Quellenlinks, Abrufdatum und unsere eigene Schrankidee.

2 CHATGPT BEDIENT DEN IN-APP-BROWSER
Öffne https://www.hornbach.de/ im In-App-Browser.
Lass ChatGPT in das Suchfeld „Möbelbauplatte Eiche“ eingeben.
Lass einen tatsächlich sichtbaren Filter anwenden. Wenn es den gewünschten Filter nicht gibt, eine verfügbare Option wählen.
Lass eine Produktseite öffnen und Preis, Einheit, Maße und Preisbasis zeigen.
Finde eine zweite Alternative. Lieferkosten und technische Unterschiede offenhalten.
Nichts bestellen oder buchen.

MÖGLICHE PRODUKTSEITEN ZUM EINSTIEG
https://www.hornbach.de/p/moebelbauplatte-eiche-sonoma-2630-x-400-x-19-mm/5520507/
https://www.hornbach.de/p/leimholzplatte-eiche-b-c-geoelt-2000-x-600-x-18-mm/8203386/
Das sind unterschiedliche Produkte, keine technisch gleichwertigen Alternativen. Preise und Verfügbarkeit im Seminar frisch prüfen. Daraus folgt keine Eignungszusage für unseren Schrank.

WENN ETWAS HÄNGT
Stopp. Welche Seite ist offen? Was ist bereits erledigt? Führe nur den nächsten Schritt aus.
Bei einer blockierten Seite einen anderen zugänglichen Baumarkt wählen und die Quellenlücke nennen.
''',
'07-weihnachtstermin.txt':'''WEIHNACHTEN MIT EINEM SCHRANK – ÜBUNG

KALENDER
25.12.2026, 10:00–10:30 Uhr, Europe/Berlin
Titel: ÜBUNG – Schrank für Hadi
Nur im eigenen Kalender. Keine Gäste, keine Einladung.
Es gibt keinen vereinbarten Termin mit Hadi Teherani.
Verschiebe denselben Testtermin auf 11:00–11:30 Uhr.
Prüfe ihn. Entferne anschließend nur diesen Übungstermin.

SPÄTER EINEN ENTWURF ERSTELLEN
Plane einen einmaligen Zeitpunkt wenige Minuten nach der Vorführung.
ChatGPT soll dann einen Mailentwurf zu unserem Schrankangebot vorbereiten.
Wenn Outlook für die geplante Aufgabe verfügbar ist: als ÜBUNG an Kerims zuvor geprüfte Adresse speichern.
Wenn Dateien oder Rechte fehlen: die Lücke melden oder den Entwurf als Text bereitstellen.
Nichts versenden. Kein zeitversetzter Mailversand.
Prüfe die gespeicherte Aufgabe, den Zeitpunkt und die Zeitzone.
Prüfe nach der Testausführung das Ergebnis und deaktiviere die Übungsaufgabe.
''',
'10-aenderung-und-uebergabe.txt':'''UNSER SCHRANKANGEBOT – ABSCHLUSSAUFTRAG

Die Arbeitszeit beträgt in unserer neuen Annahme zehn statt acht Stunden.
Keine Änderung durch Hadi Teherani: Es ist weiterhin unser Rollenspiel.

Aktualisiere die Excel-Kalkulation mit Formeln. Erhalte die ursprüngliche Version.
Fülle die Angebotsvorlage mit dem aktuellen Preis und den geprüften Angaben.
Erzeuge DOCX und PDF und lass ChatGPT sie in deinem OneDrive-Ordner speichern.
Aktualisiere den Übungsentwurf an Kerims zuvor geprüfte Adresse und Mission_Hadi.pdf als PDF-Anhang. Nicht versenden.
Zeige jemandem Quellen, Annahmen, Kalkulation, Dateien und offenen nächsten Schritt.
Räume Übungstermin und geplante Aufgabe wieder auf.

Kontrollwerte nur bei unveränderten fiktiven Startpreisen:
8 Stunden: Kosten 1.080,00 EUR netto; mit 20 % Aufschlag 1.296,00 EUR netto; mit 19 % Übungs-USt 1.542,24 EUR brutto.
10 Stunden: Kosten 1.210,00 EUR netto; mit 20 % Aufschlag 1.452,00 EUR netto; mit 19 % Übungs-USt 1.727,88 EUR brutto.
Recherchierte Materialpreise können diese Summen ändern.
'''}
for n,t in texts.items():(out/n).write_text(t)
d=Document();sec=d.sections[0];sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
style=d.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10);style.paragraph_format.space_after=Pt(5)
for n in ['Title','Heading 1','Heading 2']:d.styles[n].font.color.rgb=RGBColor(0,0,0)
d.add_paragraph('Angebot für einen Büroschrank',style='Title')
d.add_paragraph('Werkraum · Fiktiver Handwerksbetrieb · Unverbindlicher Seminarentwurf')
d.add_paragraph('Diese Vorlage wird mit Schrankbrief, Recherche und Kalkulation ausgefüllt. Es liegt keine Anfrage oder Bestellung von Hadi Teherani vor. Unbekannte Angaben bleiben offen.')
t=d.add_table(rows=0,cols=2);t.style='Light Shading Accent 1'
for label,value in [('Angebotsnummer','[eintragen]'),('Datum','[eintragen]'),('Empfänger und Büro','[an offizieller Quelle prüfen]'),('Anschrift','[Quelle oder offen]'),('Gegenstand','[Leistungsbeschreibung aus dem Schrankbrief]'),('Maße und Ausstattung','[unsere Übungsannahmen kennzeichnen]')]:
 c=t.add_row().cells;c[0].text=label;c[1].text=value
for row in t.rows:row.cells[0].width=Cm(4.7);row.cells[1].width=Cm(11.8)
d.add_paragraph('Kalkulation',style='Heading 2')
t=d.add_table(rows=1,cols=4);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Position','Menge','Einzelpreis netto','Gesamt netto']):c.text=v
for label in ['Material','Beschläge','Arbeitszeit','Transport']:
 for c,v in zip(t.add_row().cells,[label,'[ ]','[ ]','[ ]']):c.text=v
d.add_paragraph('Summe Kosten: [ ] EUR · Übungsaufschlag: [ ] % · Angebot netto: [ ] EUR\nUSt als Übungsannahme: [ ] % · Angebot brutto: [ ] EUR')
d.add_paragraph('Offene Punkte und nächste Schritte',style='Heading 2')
d.add_paragraph('[Bedarf, Maße, Ausführung, Lieferkosten und Termin prüfen. Keine bestätigten Kundenwünsche oder Zusagen erfinden.]')
d.add_paragraph('Quellen',style='Heading 2');d.add_paragraph('[Dateinamen, öffentliche Quellenlinks und Abrufdatum ergänzen.]')
sec.footer.paragraphs[0].text='Seminarübung · Kein verbindliches Angebot · Nicht an externe Empfänger versenden'
d.save(out/'05-angebotsvorlage.docx')
print('Schrankbrief, Rechercheauftrag, Kalenderübung und Angebotsvorlage erstellt.')
